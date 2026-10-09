#!/usr/bin/env python3
"""
build-statelessness-101-pptx.py — render the STATELESSNESS 101 deck.

Thin builder mirroring build-statelessness-pptx.py (the 201 compendium
deck). It reuses every slide builder, the theme, and the helpers from
build-pptx.py — loaded via importlib because that file's name has a
hyphen — and supplies its own title / agenda / closing slides plus the
introductory SECTIONS from sections_statelessness_101.py.

This is the on-ramp deck: a gentle introduction for a C++ developer who
has not read the compendium. No demo-cue slides and no diagrams — it is
text-and-code only, deliberately shallow, with pointers forward into the
201 deck's concepts.

Output: presentation/cpp-statelessness-101.pptx
"""
import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"


def load_build_pptx():
    """Load build-pptx.py (hyphenated filename) as a module."""
    spec = importlib.util.spec_from_file_location(
        "build_pptx", TOOLS / "build-pptx.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


bp = load_build_pptx()
C, F, FontFam = bp.C, bp.F, bp.FontFam
SLIDE_W, SLIDE_H = bp.SLIDE_W, bp.SLIDE_H

OUT_PATH = ROOT / "presentation" / "cpp-statelessness-101.pptx"
REPO_URL = "github.com/patterncatalyst/cpp-container-optimization-tutorial"


def build_title(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)

    dot_y = Inches(1.5)
    for i, color in enumerate([C.ACCENT_GREEN, C.ACCENT_CYAN, C.ACCENT_BLUE]):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.0 + i * 0.35), dot_y, Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    bp.add_text(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
                "STATELESSNESS 101", size=Pt(20), color=C.ACCENT_GREEN,
                bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.5),
                "A Gentle Start for C++ Services\nUnder an Orchestrator",
                size=F.TITLE_HUGE, color=C.TEXT_WHITE, bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.5),
                "Why it matters  •  kill / replace / replicate  •  "
                "request / process / external scope  •  RAII",
                size=F.BODY, color=C.TEXT_LIGHT, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.2), Inches(11.3), Inches(0.4),
                "The on-ramp to the Statelessness Compendium — "
                "no prior reading required",
                size=F.BODY_SMALL, color=C.TEXT_MUTED, italic=True, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.4),
                "Start here if you haven't read the compendium yet",
                size=F.BODY_SMALL, color=C.ACCENT_GREEN, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)

    bp.set_notes(slide,
        "Welcome. This is Statelessness 101 — the on-ramp. If you've already "
        "read the compendium, this will feel familiar and brief; it's meant "
        "for a C++ developer coming to the subject cold.\n\n"
        "Four things in twenty-some slides: why statelessness is worth "
        "caring about as a C++ developer specifically, what an orchestrator "
        "actually does to your process (kill, replace, replicate), a shallow "
        "pass over the three scopes of state, and RAII as the one mechanism "
        "you need today.\n\n"
        "This deck carries no demos of its own. Every pattern it touches has "
        "a runnable, host-verified demo in the compendium — the pointer "
        "forward is explicit at the end.")


def build_agenda(prs, total, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.TEXT_WHITE)
    bp.add_header_bar(slide, "Overview", 2, total, "Agenda — the on-ramp")
    bp.add_footer(slide, "Agenda", 2, total)

    palette = [C.ACCENT_RED, C.ACCENT_ORANGE, C.ACCENT_BLUE, C.ACCENT_GREEN,
               C.ACCENT_CYAN]
    rows_per_col = 6
    col_width = Inches(6.0)
    row_height = Inches(0.62)
    grid_left = Inches(0.5)
    grid_top = Inches(1.1)

    for i, sec in enumerate(sections):
        color = palette[i % len(palette)]
        col = i // rows_per_col
        row = i % rows_per_col
        x = grid_left + col * (col_width + Inches(0.3))
        y = grid_top + row * row_height

        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        bp.add_text(slide, x, y, Inches(0.5), Inches(0.5), str(sec["num"]),
                    size=Pt(16), color=C.TEXT_WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=bp.MSO_ANCHOR.MIDDLE,
                    font=FontFam.HEADER)
        bp.add_text(slide, x + Inches(0.65), y, Inches(5.2), Inches(0.5),
                    sec["title"], size=Pt(15), color=C.TEXT_DARK,
                    anchor=bp.MSO_ANCHOR.MIDDLE, font=FontFam.BODY)

    bp.set_notes(slide,
        "Five short sections. The first establishes why this matters for a "
        "C++ developer specifically. The second and third give you the "
        "orchestrator's vocabulary and the scope vocabulary, both at shallow "
        "depth. The fourth lands RAII as the one mechanism to take away "
        "today. The fifth is a map of where the compendium picks each "
        "thread back up.\n\n"
        "There are no demo cues in this deck — it's an introduction, not a "
        "worked example. The compendium is where you'll run things.")


def build_closing(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)
    bp.add_text(slide, Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.0),
                "Now read the compendium.", size=F.TITLE_LARGE, color=C.TEXT_WHITE,
                bold=True, font=FontFam.HEADER)
    bp.add_multi_text(slide, Inches(1.0), Inches(2.8), Inches(11.3), Inches(2.6), [
        dict(text="Statelessness is earned by where a service keeps its "
                  "data, not by how it's written.", size=Pt(18), color=C.TEXT_LIGHT,
             font=FontFam.BODY),
        dict(text="RAII makes request scope enforceable. The compendium "
                  "builds process scope and external state out just as far.",
             size=Pt(18), color=C.TEXT_LIGHT, font=FontFam.BODY, space_before=Pt(10)),
        dict(text="Eleven documents, every pattern with a runnable, "
                  "host-verified Podman demo.", size=Pt(18),
             color=C.ACCENT_GREEN, font=FontFam.BODY, space_before=Pt(10)),
    ])
    bp.add_text(slide, Inches(1.0), Inches(5.6), Inches(11.3), Inches(0.5),
                "Read the compendium  •  run the demos  •  the source is on the repo",
                size=F.BODY, color=C.TEXT_MUTED, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)
    bp.set_notes(slide,
        "That's the on-ramp. You now have the vocabulary — kill, replace, "
        "replicate; request, process, and external scope — and the one "
        "mechanism, RAII, that makes the request-scope part of it real in "
        "C++.\n\n"
        "The compendium is the next stop: the same vocabulary, developed "
        "across eleven documents, ending in a gRPC service that composes "
        "every pattern we only named today. Every one of those patterns has "
        "a runnable demo, host-verified end to end. Thanks for starting "
        "here.")


def main():
    sys.path.insert(0, str(TOOLS))
    from sections_statelessness_101 import SECTIONS

    total = 2  # title + agenda
    for s in SECTIONS:
        total += 1 + len(s["slides"])
    total += 1  # closing
    print(f"Total slides planned: {total}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    page = 1
    print(f"  [{page:3d}/{total}] Title"); build_title(prs, total); page += 1
    print(f"  [{page:3d}/{total}] Agenda"); build_agenda(prs, total, SECTIONS); page += 1

    for section in SECTIONS:
        print(f"  [{page:3d}/{total}] §{section['num']} divider — {section['title']}")
        bp.build_section_divider(prs, section["num"], section["title"],
                                 section["tagline"], page, total,
                                 section["divider_notes"])
        page += 1
        for slide_data in section["slides"]:
            kind = slide_data["kind"]
            t = slide_data.get("title") or slide_data.get("demo_name", "(demo)")
            print(f"  [{page:3d}/{total}] §{section['num']} {kind} — {t[:48]}")
            bp.render_slide(prs, section, slide_data, page, total)
            page += 1

    print(f"  [{page:3d}/{total}] Closing"); build_closing(prs, total); page += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    print(f"\nSaved: {OUT_PATH}")
    print(f"Slide count: {page - 1}")


if __name__ == "__main__":
    main()
